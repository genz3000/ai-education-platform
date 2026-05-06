"""Document Processor - Handle various file formats"""
from typing import Dict, List, Optional
import re
import json
from pathlib import Path
from io import BytesIO
from dataclasses import dataclass

from src.utils.logger import Logger

logger = Logger(__name__)


@dataclass
class ProcessedDocument:
    """Processed document result"""
    content: str
    metadata: Dict
    file_type: str
    page_count: Optional[int] = None
    language: str = "en"


class DocumentProcessor:
    """Process various document formats"""
    
    SUPPORTED_FORMATS = {
        ".txt": "text",
        ".md": "markdown",
        ".pdf": "pdf",
        ".docx": "docx",
        ".html": "html",
        ".json": "json"
    }
    
    def __init__(self):
        self.chunk_size = 1000
        self.chunk_overlap = 200
    
    async def process_file(
        self,
        file_content: bytes,
        filename: str,
        metadata: Optional[Dict] = None
    ) -> ProcessedDocument:
        """Process a file and extract content"""
        ext = Path(filename).suffix.lower()
        file_type = self.SUPPORTED_FORMATS.get(ext, "unknown")
        
        if file_type == "unknown":
            logger.warning(f"Unsupported file type: {ext}")
            raise ValueError(f"Unsupported file type: {ext}")
        
        # Extract content based on file type
        if file_type == "text":
            content = self._process_text(file_content)
        elif file_type == "markdown":
            content = self._process_markdown(file_content)
        elif file_type == "pdf":
            content = await self._process_pdf(file_content)
        elif file_type == "docx":
            content = await self._process_docx(file_content)
        elif file_type == "html":
            content = self._process_html(file_content)
        elif file_type == "json":
            content = self._process_json(file_content)
        else:
            content = file_content.decode("utf-8", errors="ignore")
        
        # Detect language (simple heuristic)
        language = self._detect_language(content)
        
        result_metadata = metadata or {}
        result_metadata.update({
            "filename": filename,
            "file_type": file_type,
            "original_size": len(file_content),
            "processed_length": len(content)
        })
        
        return ProcessedDocument(
            content=content,
            metadata=result_metadata,
            file_type=file_type,
            language=language
        )
    
    def _process_text(self, content: bytes) -> str:
        """Process plain text"""
        return content.decode("utf-8", errors="ignore")
    
    def _process_markdown(self, content: bytes) -> str:
        """Process markdown - remove formatting, keep content"""
        text = content.decode("utf-8", errors="ignore")
        # Remove markdown formatting
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)  # Links
        text = re.sub(r'[#*_`~]+', '', text)  # Headers, bold, etc
        return text.strip()
    
    async def _process_pdf(self, content: bytes) -> str:
        """Process PDF - extract text"""
        try:
            # Try using PyPDF2
            from PyPDF2 import PdfReader
            
            pdf_file = BytesIO(content)
            reader = PdfReader(pdf_file)
            
            text_parts = []
            for page in reader.pages:
                text_parts.append(page.extract_text())
            
            return "\n\n".join(text_parts)
        except ImportError:
            logger.warning("PyPDF2 not available, returning raw bytes")
            return content.decode("utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"PDF processing error: {e}")
            return content.decode("utf-8", errors="ignore")
    
    async def _process_docx(self, content: bytes) -> str:
        """Process DOCX - extract text"""
        try:
            from docx import Document
            
            doc_file = BytesIO(content)
            doc = Document(doc_file)
            
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
            
            return "\n".join(text_parts)
        except ImportError:
            logger.warning("python-docx not available")
            return content.decode("utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"DOCX processing error: {e}")
            return content.decode("utf-8", errors="ignore")
    
    def _process_html(self, content: bytes) -> str:
        """Process HTML - extract text"""
        text = content.decode("utf-8", errors="ignore")
        
        # Remove scripts and styles
        text = re.sub(r'<script[^>]*>.*?</script>', '', text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r'<style[^>]*>.*?</style>', '', text, flags=re.DOTALL | re.IGNORECASE)
        
        # Convert HTML entities
        text = text.replace('&nbsp;', ' ')
        text = text.replace('&amp;', '&')
        text = text.replace('&lt;', '<')
        text = text.replace('&gt;', '>')
        
        # Remove HTML tags
        text = re.sub(r'<[^>]+>', ' ', text)
        
        # Clean whitespace
        text = re.sub(r'\s+', ' ', text)
        
        return text.strip()
    
    def _process_json(self, content: bytes) -> str:
        """Process JSON - flatten structure"""
        try:
            data = json.loads(content)
            return self._flatten_json(data)
        except:
            return content.decode("utf-8", errors="ignore")
    
    def _flatten_json(self, data: Dict, prefix: str = "") -> str:
        """Flatten nested JSON to text"""
        parts = []
        
        if isinstance(data, dict):
            for key, value in data.items():
                new_key = f"{prefix}.{key}" if prefix else key
                if isinstance(value, (dict, list)):
                    parts.append(self._flatten_json(value, new_key))
                else:
                    parts.append(f"{key}: {value}")
        elif isinstance(data, list):
            for i, item in enumerate(data):
                parts.append(self._flatten_json(item, f"{prefix}[{i}]"))
        else:
            parts.append(f"{prefix}: {data}")
        
        return "\n".join(parts)
    
    def _detect_language(self, text: str) -> str:
        """Simple language detection"""
        # Check for Chinese characters
        chinese_chars = len(re.findall(r'[\u4e00-\u9fff]', text))
        
        if chinese_chars > len(text) * 0.3:
            return "zh"
        
        # Check for Japanese
        japanese_chars = len(re.findall(r'[\u3040-\u309f\u30a0-\u30ff]', text))
        if japanese_chars > len(text) * 0.2:
            return "ja"
        
        return "en"


class BatchDocumentProcessor:
    """Process multiple documents in batch"""
    
    def __init__(self):
        self.processor = DocumentProcessor()
    
    async def process_batch(
        self,
        files: List[Dict],
        batch_size: int = 10
    ) -> List[ProcessedDocument]:
        """Process multiple files in batches"""
        results = []
        
        for i in range(0, len(files), batch_size):
            batch = files[i:i + batch_size]
            
            for file_data in batch:
                try:
                    doc = await self.processor.process_file(
                        file_content=file_data["content"],
                        filename=file_data["name"],
                        metadata=file_data.get("metadata")
                    )
                    results.append(doc)
                except Exception as e:
                    logger.error(f"Error processing {file_data.get('name', 'unknown')}: {e}")
        
        return results


# Factory function
def create_document_processor() -> DocumentProcessor:
    """Create document processor"""
    return DocumentProcessor()